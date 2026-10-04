# ported from: src/documents/render-xlsx.test.ts
"""Test names are the snake_case form of ``describe_it``; the original Vietnamese title is the docstring.

The original unzipped the file and read the XML (``readZipEntryText``). The assertions about the XML itself
(the ``<f>`` / ``<v>`` pair, no "=" inside ``<f>``, ``ySplit``) stay on the raw part read with ``zipfile``;
the assertions about look and content (fonts, fills, number formats, merges, rich text) re-open the produced
bytes with openpyxl, because openpyxl writes its own styles table and shared-string layout (it stores strings
inline), so searching those parts for literal text would test the library, not the renderer.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from pema.documents.document_content_schema import ParseOk, Sheet, parse_sheets
from pema.documents.render_xlsx import render_xlsx, safe_sheet_name
from pema.documents.xlsx_themes import XLSX_THEMES


def _sheets(raw: list[dict[str, Any]]) -> list[Sheet]:
    result = parse_sheets(raw)
    assert isinstance(result, ParseOk), result
    return result.value


def _part(data: bytes, name: str) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read(name).decode("utf-8")


def _entries(data: bytes) -> list[str]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.namelist()


def _sheet_xml(raw: list[dict[str, Any]], index: int = 1) -> str:
    return _part(render_xlsx(_sheets(raw)), f"xl/worksheets/sheet{index}.xml")


def _open(data: bytes, *, rich_text: bool = False, data_only: bool = False) -> Workbook:
    return load_workbook(io.BytesIO(data), rich_text=rich_text, data_only=data_only)


def _ws(wb: Workbook, index: int = 0) -> Worksheet:
    sheet = wb.worksheets[index]
    assert isinstance(sheet, Worksheet)
    return sheet


def _cell(ws: Worksheet, ref: str) -> Cell:
    cell = ws[ref]
    assert isinstance(cell, Cell)
    return cell


# openpyxl's style objects are descriptors the stubs cannot type; ``Any`` is confined to these three readers.
def _font(cell: Cell) -> Any:
    return cell.font


def _fill(cell: Cell) -> Any:
    return cell.fill


def _alignment(cell: Cell) -> Any:
    return cell.alignment


def _text(value: str) -> dict[str, Any]:
    return {"kind": "text", "value": value}


def _number(value: float, fmt: str = "plain") -> dict[str, Any]:
    return {"kind": "number", "value": value, "format": fmt}


def _multiply(a: str, b: str) -> dict[str, Any]:
    return {"kind": "formula", "op": "multiply", "columns": [a, b]}


# Quote sheet: 2 goods rows (amount = qty x unit price) + 1 total row
BAO_GIA: dict[str, Any] = {
    "name": "Báo giá",
    "headers": ["Sản phẩm", "Số lượng", "Đơn giá (VND)", "Thành tiền (VND)"],
    "rows": [
        [_text("Bàn phím cơ"), _number(2), _number(1_500_000, "money"), _multiply("B", "C")],
        [_text("Chuột không dây"), _number(3), _number(450_000, "money"), _multiply("B", "C")],
        [
            _text("Tổng cộng"),
            _text(""),
            _text(""),
            {"kind": "formula", "op": "sum", "column": "D", "fromRow": 2, "toRow": 3},
        ],
    ],
}


def formula_cells(xml: str) -> list[str]:
    """Pull the formula cells out to look at ``<f>`` and ``<v>``."""
    return re.findall(r"<c[^>]*>(?:(?!</c>).)*<f>.*?</c>", xml, re.DOTALL)


def test_render_xlsx_creates_a_valid_xlsx() -> None:
    """tạo file .xlsx hợp lệ"""
    buf = render_xlsx(_sheets([BAO_GIA]))
    assert buf[:2] == b"PK"
    entries = _entries(buf)
    assert "xl/worksheets/sheet1.xml" in entries
    assert "xl/styles.xml" in entries


def test_render_xlsx_every_formula_cell_has_a_cached_v_value_the_reason_exceljs_exists() -> None:
    """MỌI ô công thức đều có giá trị cache <v> - lý do tồn tại của exceljs"""
    xml = _sheet_xml([BAO_GIA])
    cells = formula_cells(xml)
    assert len(cells) == 3, "phải có đúng 3 ô công thức"
    missing = [c for c in cells if not re.search(r"<v>", c)]
    assert missing == [], "ô công thức thiếu <v> sẽ hiện TRỐNG trong mọi công cụ xem trước"


def test_render_xlsx_cached_values_are_computed_correctly_product_and_sum() -> None:
    """giá trị cache tính ĐÚNG (nhân ra tích, SUM ra tổng)"""
    xml = _sheet_xml([BAO_GIA])
    # 2 x 1,500,000 = 3,000,000 - 3 x 450,000 = 1,350,000 - total = 4,350,000
    for expected in ["3000000", "1350000", "4350000"]:
        assert f"<v>{expected}</v>" in xml, f"thiếu giá trị {expected}"


def test_render_xlsx_cached_values_are_what_a_previewer_reads() -> None:
    """EXTRA: ``data_only=True`` is what every preview tool does - it must see the numbers, not empty cells"""
    ws = _ws(_open(render_xlsx(_sheets([BAO_GIA])), data_only=True))
    assert [_cell(ws, ref).value for ref in ("D2", "D3", "D4")] == [3_000_000, 1_350_000, 4_350_000]


def test_render_xlsx_formula_has_no_equals_sign_inside_f_ecma_376() -> None:
    """công thức KHÔNG kèm dấu = trong <f> (chuẩn ECMA-376)"""
    xml = _sheet_xml([BAO_GIA])
    assert not re.search(r"<f>=", xml), "dấu = trong <f> là sai chuẩn OOXML"
    assert "<f>B2*C2</f>" in xml, "công thức nhân phải trỏ đúng dòng"
    assert "<f>SUM(D2:D3)</f>" in xml, "công thức tổng phải trỏ đúng dải"


def test_render_xlsx_keeps_vietnamese_diacritics() -> None:
    """giữ nguyên tiếng Việt có dấu"""
    ws = _ws(_open(render_xlsx(_sheets([BAO_GIA]))))
    assert _cell(ws, "A2").value == "Bàn phím cơ"
    assert _cell(ws, "C1").value == "Đơn giá (VND)"
    assert ws.title == "Báo giá"


def test_render_xlsx_bold_header_frozen_row_1_arial_font_and_thousands_separator() -> None:
    """header đậm + đóng băng dòng 1, font Arial, số tiền có phân cách nghìn"""
    buf = render_xlsx(_sheets([BAO_GIA]))
    xml = _part(buf, "xl/worksheets/sheet1.xml")
    ws = _ws(_open(buf))
    assert "pane" in xml, "phải đóng băng dòng header"
    assert ws.freeze_panes == "A2"
    header = _cell(ws, "A1")
    assert _font(header).name == "Arial", "font phải là Arial"
    assert _font(header).b is True, "header phải in đậm"
    # "#,##0" is the built-in Excel format (numFmtId 3), so it is not declared as a string in <numFmts>
    assert _cell(ws, "C2").number_format == "#,##0", "số tiền phải áp định dạng phân cách nghìn"
    assert _cell(ws, "D2").number_format == "#,##0"


def test_render_xlsx_percent_uses_its_own_custom_format_not_the_money_one() -> None:
    """phần trăm dùng định dạng riêng (custom), không lẫn với tiền"""
    buf = render_xlsx(
        _sheets(
            [
                {
                    "name": "Tỉ lệ",
                    "headers": ["Mục", "Tỉ lệ"],
                    "rows": [[_text("Chiết khấu"), _number(0.15, "percent")]],
                }
            ]
        )
    )
    # "0.0%" is NOT a built-in format, so it must be declared in <numFmts>
    assert "0.0%" in _part(buf, "xl/styles.xml")
    assert _cell(_ws(_open(buf)), "B2").number_format == "0.0%"


def test_render_xlsx_several_sheets_are_all_created() -> None:
    """nhiều sheet đều được tạo"""
    buf = render_xlsx(
        _sheets([BAO_GIA, {"name": "Ghi chú", "headers": ["Nội dung"], "rows": [[_text("abc")]]}])
    )
    entries = _entries(buf)
    assert "xl/worksheets/sheet1.xml" in entries
    assert "xl/worksheets/sheet2.xml" in entries
    assert _open(buf).sheetnames == ["Báo giá", "Ghi chú"]


def test_render_xlsx_sheet_name_with_forbidden_characters_is_cleaned_and_the_file_stays_valid() -> None:
    """tên sheet có ký tự Excel cấm bị làm sạch, file vẫn hợp lệ"""
    assert safe_sheet_name("Báo giá: Q3/2026 [bản 1]", "Sheet1") == "Báo giá  Q3 2026  bản 1"
    assert safe_sheet_name("///", "Sheet1") == "Sheet1", "rỗng sau khi lọc thì dùng tên dự phòng"
    assert len(safe_sheet_name("x" * 50, "Sheet1")) == 31, "Excel giới hạn 31 ký tự"

    buf = render_xlsx(_sheets([{**BAO_GIA, "name": "Báo giá: Q3/2026"}]))
    assert buf[:2] == b"PK"
    assert _open(buf).sheetnames == ["Báo giá  Q3 2026"]


def test_render_xlsx_bold_marker_in_a_text_cell_becomes_bold_rich_text_without_asterisks() -> None:
    """marker **đậm** trong ô chữ thành rich text bold, dấu sao không lọt vào file"""
    buf = render_xlsx(
        _sheets(
            [
                {
                    "name": "Đậm",
                    "headers": ["**Mục**", "Giá trị"],
                    "rows": [[_text("**Tổng cộng** (đã gồm VAT)"), _number(100)]],
                }
            ]
        )
    )
    assert "**" not in _part(buf, "xl/worksheets/sheet1.xml"), "dấu ** không được xuất hiện trong file"
    ws = _ws(_open(buf, rich_text=True))
    assert _cell(ws, "A1").value == "Mục", "header phải được lọc marker"
    rich = _cell(ws, "A2").value
    assert isinstance(rich, CellRichText)
    assert "**" not in str(rich)
    assert str(rich) == "Tổng cộng (đã gồm VAT)", "chữ trong marker phải còn nguyên"
    # Rich text: the "Tổng cộng" part sits in a run that is bold
    first = rich[0]
    assert isinstance(first, TextBlock)
    assert first.text == "Tổng cộng"
    assert first.font.b is True, "phần trong ** phải in đậm"
    second = rich[1]
    assert isinstance(second, TextBlock)
    assert second.font.b in (False, None)


def test_render_xlsx_title_banner_and_subtitle_merge_across_navy_white_16pt() -> None:
    """banner title + subtitle: merge hết bề ngang, navy, chữ trắng 16pt"""
    buf = render_xlsx(
        _sheets([{**BAO_GIA, "title": "BÁO CÁO TỔNG HỢP KIM CƯƠNG", "subtitle": "Cập nhật đến 26/07/2026"}])
    )
    ws = _ws(_open(buf))
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert "A1:D1" in merged, "title phải merge hết 4 cột"
    assert "A2:D2" in merged, "subtitle phải merge hết 4 cột"
    banner = _cell(ws, "A1")
    assert _fill(banner).fgColor.rgb == "FF1F3864", "phải có màu navy"
    assert _font(banner).sz == 16, "title phải 16pt"
    assert _font(banner).color.rgb == "FFFFFFFF"
    assert banner.value == "BÁO CÁO TỔNG HỢP KIM CƯƠNG"


def test_render_xlsx_with_a_banner_formulas_shift_rows_because_the_model_counts_header_as_row_1() -> None:
    """CÓ banner thì công thức TỰ DỊCH DÒNG - model đánh số coi header là dòng 1"""
    buf = render_xlsx(_sheets([{**BAO_GIA, "title": "BÁO GIÁ", "subtitle": "Tháng 7"}]))
    sheet = _part(buf, "xl/worksheets/sheet1.xml")
    # Banner(1) + subtitle(2) + breathing row(3) + header(4) -> data from row 5
    assert "<f>B5*C5</f>" in sheet, "công thức nhân phải trỏ dòng thật (5)"
    assert "<f>SUM(D5:D6)</f>" in sheet, "SUM(D2:D3) của model phải dịch thành SUM(D5:D6)"
    # The cached values stay right whatever the row shift
    for expected in ["3000000", "1350000", "4350000"]:
        assert f"<v>{expected}</v>" in sheet, f"thiếu giá trị {expected}"
    # Freeze right below the header (row 4)
    assert re.search(r'ySplit="4"', sheet)
    assert _ws(_open(buf)).freeze_panes == "A5"


def test_render_xlsx_note_is_red_italic_text_wraps_and_hand_typed_numbers_take_the_theme_colour() -> None:
    """note cuối bảng: đỏ, nghiêng; ô chữ wrap; số nhập tay tô màu theo theme"""
    note = "Lưu ý: số liệu là cáo buộc ban đầu, chưa phải bản án."
    buf = render_xlsx(_sheets([{**BAO_GIA, "theme": "burgundy", "note": note}]))
    ws = _ws(_open(buf))
    # header row 1, data rows 2-4, breathing row 5, note row 6
    note_cell = _cell(ws, "A6")
    assert note_cell.value == note, "note phải vào file"
    assert _font(note_cell).color.rgb == "FFC00000", "note phải màu đỏ cảnh báo (không đổi theo theme)"
    assert _font(note_cell).i is True, "note phải in nghiêng"
    # Hand-typed numbers take the theme colour for a harmonious page; formula cells stay black -
    # the reader tells original data from computed results (xlsx skill convention)
    assert _font(_cell(ws, "C2")).color.rgb == "FF8B2635", "số nhập tay phải mang màu của theme đang dùng"
    assert _font(_cell(ws, "D2")).color is None, "ô công thức để đen"
    assert _alignment(_cell(ws, "A2")).wrap_text is True, "ô chữ phải wrap"


def test_render_xlsx_themes_are_switchable_each_tone_gives_a_different_file() -> None:
    """theme đổi được: mỗi tông ra file màu KHÁC nhau (10 báo cáo không ra 10 file giống hệt)"""

    def banner_fill(theme: str | None) -> str:
        sheet: dict[str, Any] = {**BAO_GIA, "title": "BÁO CÁO"}
        if theme is not None:
            sheet["theme"] = theme
        ws = _ws(_open(render_xlsx(_sheets([sheet]))))
        return str(_fill(_cell(ws, "A1")).fgColor.rgb)

    navy, green, burgundy = banner_fill("navy"), banner_fill("green"), banner_fill("burgundy")
    assert navy == "FF1F3864", "navy phải dùng màu navy"
    assert green == "FF1E6B3A", "green phải dùng màu xanh lá"
    assert burgundy == "FF8B2635", "burgundy phải dùng màu đỏ rượu"
    assert len({navy, green, burgundy}) == 3, "hai theme phải cho ra style khác nhau"
    # No theme chosen -> navy by default
    assert banner_fill(None) == "FF1F3864"


def test_render_xlsx_every_theme_meets_wcag_aa_for_white_text_on_the_header_background() -> None:
    """mọi theme đều đạt WCAG AA (>=4.5:1) cho chữ trắng trên nền header"""

    def luminance(hex_color: str) -> float:
        r, g, b = (
            (lambda c: c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4)(
                int(hex_color[i : i + 2], 16) / 255
            )
            for i in (0, 2, 4)
        )
        return 0.2126 * r + 0.7152 * g + 0.0722 * b

    for name, theme in XLSX_THEMES.items():
        # Drop the first 2 alpha characters of the ARGB
        bg = luminance(theme.header[2:])
        ratio = (1 + 0.05) / (bg + 0.05)
        assert ratio >= 4.5, f'theme "{name}" chỉ đạt {ratio:.2f}:1, chữ trắng khó đọc'


def test_render_xlsx_alternating_stripe_even_rows_get_a_light_background() -> None:
    """sọc xen kẽ: dòng chẵn có nền nhạt để mắt không lạc dòng"""
    ws = _ws(_open(render_xlsx(_sheets([BAO_GIA]))))
    assert _fill(_cell(ws, "A3")).fgColor.rgb == "FFF2F5FA", "phải có nền sọc nhạt của theme navy"
    assert _fill(_cell(ws, "A2")).fill_type is None, "dòng lẻ không có nền"


def test_render_xlsx_formula_that_cannot_be_computed_still_writes_0_not_empty() -> None:
    """ô công thức không tính được vẫn ghi số 0, không để trống"""
    # Multiply by a text cell -> no number comes out
    xml = _sheet_xml(
        [
            {
                "name": "Lỗi",
                "headers": ["A", "B", "C"],
                "rows": [[_text("chữ"), _text("chữ"), _multiply("A", "B")]],
            }
        ]
    )
    cells = formula_cells(xml)
    assert len(cells) == 1
    assert re.search(r"<v>0</v>", cells[0]), "không tính được thì ghi 0 thay vì để trống"


# ------------------------------------------------------------------ EXTRA: openpyxl-specific guards


def test_extra_text_cell_starting_with_equals_stays_text_never_a_formula() -> None:
    """Formula injection guard: openpyxl would turn "=..." into a formula, exceljs kept the string."""
    payload = '=HYPERLINK("https://example.invalid","x")'
    buf = render_xlsx(
        _sheets(
            [
                {
                    "name": "S",
                    "title": payload,
                    "headers": [payload],
                    "rows": [[_text(payload)]],
                    "note": payload,
                }
            ]
        )
    )
    xml = _part(buf, "xl/worksheets/sheet1.xml")
    assert "<f>" not in xml
    ws = _ws(_open(buf))
    for ref in ("A1", "A3", "A4"):
        assert _cell(ws, ref).data_type != "f"
        assert _cell(ws, ref).value == payload


def test_extra_single_column_banner_is_not_merged_and_the_file_reopens() -> None:
    buf = render_xlsx(
        _sheets([{"name": "S", "title": "T", "headers": ["A"], "rows": [[_text("x")]], "note": "n"}])
    )
    ws = _ws(_open(buf))
    assert list(ws.merged_cells.ranges) == []
    assert _cell(ws, "A1").value == "T"


def test_extra_duplicate_sheet_names_are_numbered_instead_of_failing() -> None:
    other = {**BAO_GIA, "name": "báo giá"}
    third = {**BAO_GIA, "name": "Báo giá"}
    buf = render_xlsx(_sheets([BAO_GIA, other, third]))
    assert _open(buf).sheetnames == ["Báo giá", "báo giá 2", "Báo giá 3"]


def test_extra_cached_values_reach_every_sheet_of_a_multi_sheet_workbook() -> None:
    buf = render_xlsx(_sheets([BAO_GIA, {**BAO_GIA, "name": "Bản 2", "title": "T"}]))
    wb = _open(buf, data_only=True)
    assert _cell(_ws(wb, 0), "D4").value == 4_350_000
    assert _cell(_ws(wb, 1), "D6").value == 4_350_000


def test_extra_fractional_and_empty_cells_are_written_cleanly() -> None:
    buf = render_xlsx(
        _sheets(
            [
                {
                    "name": "S",
                    "headers": ["A", "B", "C"],
                    "rows": [[_number(2.5), _text(""), _multiply("A", "A")]],
                }
            ]
        )
    )
    xml = _part(buf, "xl/worksheets/sheet1.xml")
    assert "<v>2.5</v>" in xml
    assert "<v>6.25</v>" in xml
    assert not re.search(r't="inlineStr"\s*/>', xml), "no empty string records"
