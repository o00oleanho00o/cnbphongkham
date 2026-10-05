# ported from: src/documents/render-docx.ts
"""Build a .docx file from the content the model supplies.

Every formatting convention lives HERE and not in the prompt - so the file is always equally good even when a
cheap model runs. The typography (Times New Roman, black headings, standard margins) is in
``render_docx_styles.py`` and the tables in ``render_docx_tables.py``; the structural gotchas come from
Anthropic's ``docx`` skill and were verified by unzipping the file and reading the XML back.

Forced deviations: ``docx`` -> ``python-docx``; ``Packer.toBuffer`` (async) -> a SYNC function returning
``bytes`` (rendering is CPU work with no I/O; the tool layer calls it as the original did, ``await`` is the
tool's concern). ``numbering.config`` became an ``abstractNum`` + ``num`` pair appended to the template's
numbering part (python-docx cannot create a numbering part, the template already has one).
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Final

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Twips

from pema.documents.document_content_schema import (
    BulletsBlock,
    DocumentBlock,
    HeadingBlock,
    ParagraphBlock,
    TableBlock,
)
from pema.documents.docx_text_runs import add_text_runs, parse_text_runs
from pema.documents.render_docx_styles import Xml, apply_docx_styles
from pema.documents.render_docx_tables import build_table, build_two_columns

BULLET_REFERENCE: Final = "bullet-list"
"""Numbering reference shared by every bullets block of the document (written as the abstractNum name)."""


def heading_cua_cap(level: int) -> str:
    """Heading level -> the built-in paragraph style name.

    A FUNCTION and not a lookup table, because the schema declares ``level`` as a number range (see
    ``document_content_schema.py`` for why it is not a union of literals). With an explicit default branch
    instead of a cast: widening the range in the schema later still yields a valid heading, not ``None``."""
    if level <= 1:
        return "Heading 1"
    if level >= 3:
        return "Heading 3"
    return "Heading 2"


_ALIGNMENT: Final = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

FIRST_LINE_INDENT: Final = 567
"""1cm first-line indent for a normal paragraph - the Vietnamese document layout habit."""


def _add_bullet_numbering(document: DocumentObject) -> int:
    """Append the bullet list definition to the numbering part; return its ``numId``."""
    numbering: Xml = document.part.numbering_part.element
    abstract_ids = [int(v) for v in numbering.xpath("w:abstractNum/@w:abstractNumId")]
    num_ids = [int(v) for v in numbering.xpath("w:num/@w:numId")]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = parse_xml(
        f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abstract_id}">'
        '<w:multiLevelType w:val="singleLevel"/>'
        f'<w:name w:val="{BULLET_REFERENCE}"/>'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="•"/>'
        '<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>'
        "</w:abstractNum>"
    )
    num = parse_xml(
        f'<w:num {nsdecls("w")} w:numId="{num_id}"><w:abstractNumId w:val="{abstract_id}"/></w:num>'
    )
    # Schema order: every abstractNum comes before the first num, and numIdMacAtCleanup (if any) stays last
    nums = numbering.findall(qn("w:num"))
    if not nums:
        numbering.append(abstract)
        numbering.append(num)
    else:
        nums[0].addprevious(abstract)
        nums[-1].addnext(num)
    return num_id


def _add_block(document: DocumentObject, block: DocumentBlock, bullet_num_id: int) -> None:
    if isinstance(block, HeadingBlock):
        # The built-in heading style (already restyled black) so a table of contents can recognise it if
        # that is ever needed
        document.add_paragraph(block.text, style=heading_cua_cap(block.level))
        return

    if isinstance(block, ParagraphBlock):
        # The model may put "\n" in - split into several paragraphs, because a line break character would
        # otherwise be ignored and the whole paragraph would stick into one long line
        align = block.align or "justify"
        for line in (raw.strip() for raw in block.text.split("\n")):
            if not line:
                continue
            paragraph = document.add_paragraph()
            add_text_runs(paragraph, parse_text_runs(line))
            paragraph.alignment = _ALIGNMENT[align]
            # Indent only for running prose; a centred/right line is a closing line (place, date, signature)
            # and indenting it misaligns it
            if align == "justify":
                paragraph.paragraph_format.first_line_indent = Twips(FIRST_LINE_INDENT)
        return

    if isinstance(block, BulletsBlock):
        # Never put a "•" character straight into the text - it must go through the numbering config,
        # otherwise Word does not recognise a list (no indent, no continuation)
        for item in block.items:
            paragraph = document.add_paragraph()
            add_text_runs(paragraph, parse_text_runs(item))
            p: Xml = paragraph._p  # pyright: ignore[reportPrivateUsage]  # python-docx has no public handle on the oxml tree
            num_pr = p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_ilvl().val = 0
            num_pr.get_or_add_numId().val = bullet_num_id
            paragraph.paragraph_format.space_after = Twips(60)
        return

    if isinstance(block, TableBlock):
        build_table(document, block.headers, block.rows)
        # An empty paragraph after a table: 2 adjacent tables without it get merged into one by Word
        document.add_paragraph("")
        return

    # The union is closed: what is left is a TwoColumnsBlock
    build_two_columns(document, block.left, block.right)
    document.add_paragraph("")


@dataclass(frozen=True)
class DocxMeta:
    title: str | None = None


def render_docx(blocks: list[DocumentBlock], meta: DocxMeta | None = None) -> bytes:
    document = Document()
    apply_docx_styles(document)
    # The template's author is "python-docx"; this file's author is nobody in particular
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""

    bullet_num_id = _add_bullet_numbering(document)

    title = meta.title.strip() if meta is not None and meta.title else ""
    if title:
        document.add_paragraph(title, style="Title")
    for block in blocks:
        _add_block(document, block, bullet_num_id)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
