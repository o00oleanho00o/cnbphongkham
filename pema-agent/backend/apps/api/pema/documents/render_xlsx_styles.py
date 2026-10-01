# ported from: src/documents/render-xlsx-styles.ts
"""How the .xlsx file looks - lifted from a sample file made by Claude (the xlsx skill) that the user rated
"beautiful": white title banner on a dark background, wrapped + bordered table header, coloured numbers, red
italic note.

COLOUR comes from the theme (the model picks it by context - see ``xlsx_themes.py``); everything else (font
size, row height, how cells merge, wrap) is fixed here so two files of the same theme always have the same
quality.

Forced deviations (exceljs -> openpyxl):
* ``ws.addRow`` appends after ``max_row`` in exceljs; openpyxl's ``append`` of an EMPTY row does not move
  ``max_row``, so the row number is explicit: each ``add_*`` function takes the ``row`` to write and
  RETURNS the next free row.
* A merge over a single column (``A1:A1``) is skipped: Excel treats a one-cell merge as a corrupt record and
  offers to repair the file; exceljs happened to write it, openpyxl would too.
* Text is always stored as TEXT (``set_text``): openpyxl turns any string that starts with "=" into a formula
  (exceljs kept strings as strings). A model-supplied "=cmd|..." in a header, title or text cell must stay
  words, never become a formula (spreadsheet formula injection).
"""

from __future__ import annotations

import math
from typing import Final

from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pema.documents.xlsx_themes import XlsxTheme

NOTE_RED: Final = "FFC00000"
SUBTITLE_GRAY: Final = "FFD9D9D9"
FONT_NAME: Final = "Arial"
FONT_SIZE: Final = 10


def cell_at(ws: Worksheet, row: int, column: int) -> Cell:
    """The writable cell at (row, column). Always taken BEFORE merging: once a range is merged, every cell but
    its top-left one becomes a read-only ``MergedCell``."""
    cell = ws.cell(row=row, column=column)
    if not isinstance(cell, Cell):
        raise TypeError(f"Cell ({row}, {column}) is part of a merged range")
    return cell


def set_text(cell: Cell, text: str) -> None:
    """Store ``text`` as a string even when it starts with "=" (see the module docstring). An empty string
    leaves the cell without a value (only its style): a string cell with no content is a malformed record."""
    if not text:
        cell.value = None
        return
    cell.value = text
    cell.data_type = "s"


def cell_border(theme: XlsxTheme) -> Border:
    thin = Side(style="thin", color=theme.border)
    return Border(top=thin, bottom=thin, left=thin, right=thin)


def _solid(argb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=argb)


def _merge_across(ws: Worksheet, row: int, column_count: int) -> None:
    if column_count > 1:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=column_count)


def add_title_row(ws: Worksheet, row: int, title: str, column_count: int, theme: XlsxTheme) -> int:
    """Title banner row: merged across the whole width, theme background, white 16pt text."""
    cell = cell_at(ws, row, 1)
    set_text(cell, title)
    # The height is exactly what was measured on the sample file (25.5 / 15.75 / 30), not a guess
    ws.row_dimensions[row].height = 25.5
    _merge_across(ws, row, column_count)
    cell.font = Font(name=FONT_NAME, size=16, bold=True, color="FFFFFFFF")
    cell.fill = _solid(theme.header)
    cell.alignment = Alignment(vertical="center")
    return row + 1


def add_subtitle_row(ws: Worksheet, row: int, subtitle: str, column_count: int, theme: XlsxTheme) -> int:
    """Subtitle row under the banner: italic, light grey on the same theme background."""
    cell = cell_at(ws, row, 1)
    set_text(cell, subtitle)
    ws.row_dimensions[row].height = 15.75
    _merge_across(ws, row, column_count)
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE, italic=True, color=SUBTITLE_GRAY)
    cell.fill = _solid(theme.header)
    cell.alignment = Alignment(vertical="center")
    return row + 1


def style_header_row(ws: Worksheet, row: int, column_count: int, theme: XlsxTheme) -> None:
    """Table header row: bold white text on the theme background, wrapped, centred, full border."""
    ws.row_dimensions[row].height = 30
    border = cell_border(theme)
    for column in range(1, column_count + 1):
        cell = cell_at(ws, row, column)
        cell.font = Font(name=FONT_NAME, size=FONT_SIZE, bold=True, color="FFFFFFFF")
        cell.fill = _solid(theme.header)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = border


def add_note_row(ws: Worksheet, row: int, note: str, column_count: int) -> int:
    """Note row at the end of the table: red, italic, merged across, wrapped. One blank row first."""
    row += 1  # the blank row (ws.addRow([]) in the original)
    cell = cell_at(ws, row, 1)
    set_text(cell, note)
    _merge_across(ws, row, column_count)
    cell.font = Font(name=FONT_NAME, size=FONT_SIZE, italic=True, color=NOTE_RED)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    # A long note needs room to breathe - estimate the number of lines from the total width
    total_width = sum(
        ws.column_dimensions[get_column_letter(column)].width or 10 for column in range(1, column_count + 1)
    )
    ws.row_dimensions[row].height = max(16.0, math.ceil(len(note) / max(20.0, total_width * 0.9)) * 14.0)
    return row + 1
