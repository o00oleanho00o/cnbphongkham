# ported from: src/documents/render-xlsx.ts
"""Build an .xlsx file from the content the model supplies.

IRON RULE: every formula cell must carry its ``result``. Without it the file has only ``<f>`` and no ``<v>``,
and every preview tool shows an EMPTY cell. That was the one reason the project chose exceljs over a lighter
library - see the regression tests in test_render_xlsx.py.

Style (navy banner, white header on navy, blue numbers, red note) lifted from the sample file of Anthropic's
xlsx skill - it lives in ``render_xlsx_styles.py``.

Forced deviations (exceljs -> openpyxl):
* openpyxl writes a formula as ``<f>..</f><v/>`` and has NO way to attach the computed result, which is
  exactly what the iron rule needs. So the workbook is written first, remembering the computed value of
  every formula
  cell (``computeSheetValues``, unchanged), and then ``_inject_cached_values`` rewrites each ``<v>`` in the
  worksheet XML of the produced package. Nothing else in the package is touched.
* ``Buffer`` -> ``bytes``; ``renderXlsx`` was async only because of ``writeBuffer`` - it is a SYNC function
  here (CPU work, no I/O).
* ``ws.addRow`` -> explicit row numbers (see ``render_xlsx_styles.py``); a workbook with no sheet is refused
  (``ValueError``) because openpyxl cannot save one - the tool schema already demands at least 1 sheet.
* Rich text: ``CellRichText`` / ``TextBlock`` / ``InlineFont`` replace the ``richText`` array of exceljs.
"""

from __future__ import annotations

import io
import re
import zipfile
from typing import Final
from xml.etree import ElementTree as ET  # only used on the package openpyxl has just produced
from xml.etree.ElementTree import Element

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from pema.documents.document_content_schema import (
    MultiplyFormulaCell,
    NumberCell,
    Sheet,
    SpreadsheetCellOutput,
    TextCell,
)
from pema.documents.render_xlsx_styles import (
    FONT_NAME,
    FONT_SIZE,
    add_note_row,
    add_subtitle_row,
    add_title_row,
    cell_at,
    cell_border,
    set_text,
    style_header_row,
)
from pema.documents.spreadsheet_formula_values import compute_sheet_values
from pema.documents.xlsx_themes import XlsxTheme, resolve_theme

# Measured on the sample file rated "beautiful": STT column 5, content column 34-42.
# Wider than 42 and a long line is too long for the eye, narrower than 5 and a 2-digit number becomes ###
MIN_COLUMN_WIDTH: Final = 5
MAX_COLUMN_WIDTH: Final = 42

NUMBER_FORMAT: Final = {
    "plain": "#,##0",
    "money": "#,##0",
    # Stored as a fraction: 0.15 shows as 15.0%
    "percent": "0.0%",
}

_BOLD_SPLIT_RE: Final = re.compile(r"\*\*([^*]+)\*\*")
_INVALID_SHEET_CHARS_RE: Final = re.compile(r"[:\\/?*\[\]]")


def safe_sheet_name(raw: str, fallback: str) -> str:
    """Excel forbids : \\ / ? * [ ] in a sheet name and limits it to 31 characters."""
    cleaned = _INVALID_SHEET_CHARS_RE.sub(" ", raw).strip()[:31]
    return cleaned or fallback


def _text_cell_value(text: str) -> str | CellRichText:
    """A text cell with ``**bold**`` markers -> rich text. The model is taught this marker (the Word tool) so
    the Excel cell must understand it too - without parsing, the asterisks land in the file as they are. The
    font must be restated on every run because rich text overrides the font of the cell."""
    if "**" not in text:
        return text
    parts = _BOLD_SPLIT_RE.split(text)
    blocks = [
        TextBlock(InlineFont(rFont=FONT_NAME, sz=FONT_SIZE, b=index % 2 == 1), part)
        for index, part in enumerate(parts)
        if part
    ]
    return CellRichText(*blocks) if blocks else ""


def _plain_number(value: float) -> float | int:
    return int(value) if value == int(value) and abs(value) < 2**53 else value


def _write_cell(
    ws: Worksheet,
    row: int,
    column: int,
    cell: SpreadsheetCellOutput,
    row_offset: int,
    computed: float | None,
    cached: dict[str, float],
) -> None:
    """Value for the sheet; a formula cell always gets its pre-computed result remembered in ``cached``
    (written into ``<v>`` afterwards).

    ``row_offset``: the model numbers rows treating the header as row 1, but the banner/subtitle push
    the table down - every row reference inside a formula must add the offset, otherwise the formula
    points into the banner."""
    target = cell_at(ws, row, column)
    if isinstance(cell, TextCell):
        value = _text_cell_value(cell.value)
        if isinstance(value, str):
            set_text(target, value)
        else:
            target.value = value
        return
    if isinstance(cell, NumberCell):
        target.value = _plain_number(cell.value)
        return

    formula = (
        f"{cell.columns[0]}{row}*{cell.columns[1]}{row}"
        if isinstance(cell, MultiplyFormulaCell)
        else f"SUM({cell.column}{cell.from_row + row_offset}:{cell.column}{cell.to_row + row_offset})"
    )
    target.value = f"={formula}"
    # The result must be a number; if it cannot be computed, 0 is better than empty (an empty cell in the
    # preview looks like an error, a 0 tells the reader at once there is no data yet)
    cached[target.coordinate] = computed if computed is not None else 0.0


def _style_data_cell(
    target_ws: Worksheet,
    row: int,
    column: int,
    cell: SpreadsheetCellOutput,
    theme: XlsxTheme,
    *,
    is_stripe: bool,
) -> None:
    """Style of a data cell. By the xlsx skill's convention: hand-typed numbers are COLOURED, formula cells
    stay black - the reader tells original data from computed results. The number colour comes from the theme
    (dark, same hue) instead of pure blue so the whole page stays harmonious."""
    target = cell_at(target_ws, row, column)
    target.border = cell_border(theme)
    # Very light alternating stripe - the eye following a long table does not lose its row
    if is_stripe:
        target.fill = PatternFill(fill_type="solid", fgColor=theme.stripe)
    if isinstance(cell, TextCell):
        target.font = Font(name=FONT_NAME, size=FONT_SIZE)
        target.alignment = Alignment(wrap_text=True, vertical="top")
        return
    if isinstance(cell, NumberCell):
        target.font = Font(name=FONT_NAME, size=FONT_SIZE, color=theme.number)
        target.number_format = NUMBER_FORMAT[cell.format]
        return
    target.font = Font(name=FONT_NAME, size=FONT_SIZE)
    target.number_format = NUMBER_FORMAT["money"]


def _vi_locale_string(value: float) -> str:
    """``Number.prototype.toLocaleString("vi-VN")``: "." groups thousands, "," starts at most 3 decimals."""
    int_part, _, frac = f"{abs(value):,.3f}".partition(".")
    frac = frac.rstrip("0")
    text = int_part.replace(",", ".") + (f",{frac}" if frac else "")
    return f"-{text}" if value < 0 else text


def _column_width(header: str, rows: list[list[SpreadsheetCellOutput]], col_index: int) -> float:
    """Column width from the longest content, clamped to an easy-to-read range (longer wraps)."""
    longest = len(header)
    for row in rows:
        if col_index >= len(row):
            continue
        cell = row[col_index]
        if isinstance(cell, TextCell):
            text = cell.value
        elif isinstance(cell, NumberCell):
            text = _vi_locale_string(cell.value)
        else:
            text = "0.000.000"
        longest = max(longest, len(text))
    return float(min(MAX_COLUMN_WIDTH, max(MIN_COLUMN_WIDTH, longest + 2)))


# ----------------------------------------------------------------- cached values of formula cells

_MAIN_NS: Final = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS: Final = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_SHEET_PART_RE: Final = re.compile(r"xl/worksheets/sheet(\d+)\.xml")


def _format_cached(value: float) -> str:
    return str(int(value)) if value == int(value) and abs(value) < 1e15 else repr(value)


def _patch_sheet_xml(payload: bytes, cached: dict[str, float]) -> bytes:
    ET.register_namespace("", _MAIN_NS)
    ET.register_namespace("r", _REL_NS)
    root = ET.fromstring(payload)  # noqa: S314  # the XML openpyxl produced a moment ago, not untrusted input
    c_tag, f_tag, v_tag = (f"{{{_MAIN_NS}}}{name}" for name in ("c", "f", "v"))
    for c in root.iter(c_tag):
        if c.find(f_tag) is None:
            continue
        value = cached.get(c.get("r", ""))
        if value is None:
            continue
        v: Element | None = c.find(v_tag)
        if v is None:
            v = ET.SubElement(c, v_tag)
        v.text = _format_cached(value)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _inject_cached_values(data: bytes, cached_by_sheet: dict[int, dict[str, float]]) -> bytes:
    """Fill ``<v>`` of every formula cell (``cached_by_sheet``: 1-based sheet number -> coordinate -> value).
    Every other part of the package is copied byte for byte."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            payload = zin.read(item.filename)
            match = _SHEET_PART_RE.fullmatch(item.filename)
            if match and cached_by_sheet.get(int(match.group(1))):
                payload = _patch_sheet_xml(payload, cached_by_sheet[int(match.group(1))])
            zout.writestr(item, payload)
    return out.getvalue()


# ----------------------------------------------------------------- the renderer


def render_xlsx(sheets: list[Sheet]) -> bytes:
    if not sheets:
        raise ValueError("Cần ít nhất 1 sheet")
    workbook = Workbook()
    workbook.properties.creator = ""
    default_sheet = workbook.active
    if default_sheet is not None:
        workbook.remove(default_sheet)
    used_names: set[str] = set()
    cached_by_sheet: dict[int, dict[str, float]] = {}

    for sheet_index, sheet in enumerate(sheets):
        # Sheet names must be unique (also after sanitising, and case-insensitive) - number the later one
        # instead of letting the whole file die
        fallback = f"Sheet{sheet_index + 1}"
        name = safe_sheet_name(sheet.name, fallback)
        n = 2
        while name.lower() in used_names:
            name = f"{safe_sheet_name(sheet.name, fallback)[:28]} {n}"
            n += 1
        used_names.add(name.lower())
        ws = workbook.create_sheet(title=name)
        cached: dict[str, float] = {}
        cached_by_sheet[sheet_index + 1] = cached

        column_count = len(sheet.headers)
        computed_values = compute_sheet_values(sheet)
        theme = resolve_theme(sheet.theme)

        # Set the column widths BEFORE adding rows so add_note_row can estimate the height
        for index, header in enumerate(sheet.headers):
            ws.column_dimensions[get_column_letter(index + 1)].width = _column_width(
                header, sheet.rows, index
            )

        row = 1
        if sheet.title:
            row = add_title_row(ws, row, sheet.title, column_count, theme)
        if sheet.subtitle:
            row = add_subtitle_row(ws, row, sheet.subtitle, column_count, theme)
        if sheet.title or sheet.subtitle:
            row += 1  # breathing row between the banner and the table

        # The header is already bold - only drop the ** marker if the model put one in
        header_row = row
        for index, header in enumerate(sheet.headers):
            set_text(cell_at(ws, header_row, index + 1), header.replace("**", ""))
        style_header_row(ws, header_row, column_count, theme)
        # The model numbers rows treating header = row 1; a banner pushing the table down means every
        # reference in a formula has to shift along
        row_offset = header_row - 1

        row = header_row + 1
        for row_index, data_row in enumerate(sheet.rows):
            excel_row = header_row + 1 + row_index
            for col_index, cell in enumerate(data_row):
                computed_row = computed_values[row_index]
                computed = computed_row[col_index] if col_index < len(computed_row) else None
                _write_cell(ws, excel_row, col_index + 1, cell, row_offset, computed, cached)
                _style_data_cell(ws, excel_row, col_index + 1, cell, theme, is_stripe=row_index % 2 == 1)
            row = excel_row + 1

        if sheet.note:
            row = add_note_row(ws, row, sheet.note, column_count)

        # Scrolling a long table still shows the banner + the column names
        ws.freeze_panes = f"A{header_row + 1}"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return _inject_cached_values(buffer.getvalue(), cached_by_sheet)
