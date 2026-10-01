# ported from: src/documents/render-docx-tables.ts
"""Table-building part of the .docx renderer - split out of render_docx.py to keep it short.

Forced deviation: ``docx`` builds tables from option objects; ``python-docx`` adds a table to a ``Document``
and the properties the original passed as options (table width, borders, cell margins, header repeat, cell
shading) are written as raw WordprocessingML here, in schema order. ``build_table`` / ``build_two_columns``
therefore take the ``Document`` they append to.
"""

from __future__ import annotations

from typing import Final

from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Twips
from docx.table import Table

from pema.documents.docx_text_runs import add_text_runs, parse_text_runs
from pema.documents.render_docx_styles import (
    CONTENT_WIDTH_DXA,
    TABLE_CELL_MARGINS,
    Xml,
    set_child,
    xml_element,
)

# Children of w:tblPr in schema order
_TBL_PR_ORDER: Final = (
    "w:tblStyle",
    "w:tblpPr",
    "w:tblOverlap",
    "w:bidiVisual",
    "w:tblStyleRowBandSize",
    "w:tblStyleColBandSize",
    "w:tblW",
    "w:jc",
    "w:tblCellSpacing",
    "w:tblInd",
    "w:tblBorders",
    "w:shd",
    "w:tblLayout",
    "w:tblCellMar",
    "w:tblLook",
)

_BORDER_SIDES: Final = ("top", "left", "bottom", "right", "insideH", "insideV")
_NO_BORDER: Final = {"val": "none", "sz": "0", "space": "0", "color": "FFFFFF"}
_THIN_BORDER: Final = {"val": "single", "sz": "4", "space": "0", "color": "auto"}


def column_widths(count: int, total: int = CONTENT_WIDTH_DXA) -> list[int]:
    """Split the width evenly over N columns, the remainder goes to the last column so the sum matches
    exactly."""
    each = total // count
    widths = [each] * count
    widths[count - 1] = total - each * (count - 1)
    return widths


def _style_table(table: Table, widths: list[int], *, hidden_borders: bool) -> None:
    tbl: Xml = table._tbl  # pyright: ignore[reportPrivateUsage]  # python-docx has no public handle on the oxml tree
    tbl_pr: Xml = tbl.tblPr
    # Dual width: the table AND every cell, both DXA. PERCENTAGE looks fine in Word but breaks in Google Docs.
    set_child(tbl_pr, xml_element("w:tblW", w=str(CONTENT_WIDTH_DXA), type="dxa"), _TBL_PR_ORDER)

    borders = xml_element("w:tblBorders")
    spec = _NO_BORDER if hidden_borders else _THIN_BORDER
    for side in _BORDER_SIDES:
        borders.append(xml_element(f"w:{side}", **spec))
    set_child(tbl_pr, borders, _TBL_PR_ORDER)

    # Padding inside the cell - without it the text sticks to the border
    margins = xml_element("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        margins.append(xml_element(f"w:{side}", w=str(TABLE_CELL_MARGINS[side]), type="dxa"))
    set_child(tbl_pr, margins, _TBL_PR_ORDER)

    for index, width in enumerate(widths):
        table.columns[index].width = Twips(width)  # the w:gridCol


def _build_cell(cell: Xml, text: str, width: int, *, is_header: bool) -> None:
    cell.width = Twips(width)
    if is_header:
        # CLEAR and not SOLID - SOLID renders a solid black background
        tc: Xml = cell._tc  # pyright: ignore[reportPrivateUsage]  # python-docx has no public handle on the oxml tree
        tc.get_or_add_tcPr().append(xml_element("w:shd", val="clear", color="auto", fill="F1F5F9"))
    paragraph = cell.paragraphs[0]
    add_text_runs(paragraph, parse_text_runs(text, base_bold=is_header))
    paragraph.paragraph_format.space_after = Twips(0)


def build_table(document: DocumentObject, headers: list[str], rows: list[list[str]]) -> Table:
    widths = column_widths(len(headers))
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    _style_table(table, widths, hidden_borders=False)

    header_row = table.rows[0]
    for index, text in enumerate(headers):
        _build_cell(header_row.cells[index], text, widths[index], is_header=True)
    # Repeat the header when the table spills onto the next page
    tr: Xml = header_row._tr  # pyright: ignore[reportPrivateUsage]  # python-docx has no public handle on the oxml tree
    tr.get_or_add_trPr().append(xml_element("w:tblHeader"))

    for row_index, row in enumerate(rows):
        cells = table.rows[row_index + 1].cells
        for index, text in enumerate(row):
            _build_cell(cells[index], text, widths[index], is_header=False)
    return table


def build_two_columns(document: DocumentObject, left: list[str], right: list[str]) -> Table:
    """Two borderless columns side by side - the head of an administrative document (agency | national motto)
    and the signature block (recipients | position + name). Each line is centred in its column, bold through
    the ** marker the model sets itself."""
    widths = column_widths(2)
    table = document.add_table(rows=1, cols=2)
    _style_table(table, widths, hidden_borders=True)
    row_cells = table.rows[0].cells
    for cell, lines, width in ((row_cells[0], left, widths[0]), (row_cells[1], right, widths[1])):
        cell.width = Twips(width)
        for index, line in enumerate(lines):
            paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            add_text_runs(paragraph, parse_text_runs(line))
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Twips(0)
    return table
